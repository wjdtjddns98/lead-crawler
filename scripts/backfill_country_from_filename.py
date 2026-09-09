# 1회성(2026-08-21, git 미반영): 국가 빈값을 원본 엑셀 파일명/시트명 국가로 백필.
"""fastpath 인제스트가 국가 '셀' 값만 썼는데, 원본은 파일명/시트명에 국가를 나눠놨다
("대한민국 비상장사 리스트", "해외 상장사 - 싱가포르"). 그 신호로 빈 국가를 채운다.

파일명·시트명에서 국가 토큰(한글/영문)을 찾아 ISO2 로 — 못 찾으면 그 파일은 스킵
(해외 혼합 리스트 등). 현재 빈값/쓰레기값인 행만 갱신(기존 값·LLM 판정 보존).

사용: python scripts/backfill_country_from_filename.py [--apply]
"""

from __future__ import annotations

import argparse
import glob
import os
import re

from openpyxl import load_workbook
from sqlalchemy import text

from leadcrawler.config import Settings
from leadcrawler.dedup import canonical_key, normalize_domain
from leadcrawler.storage.db import get_sessionmaker

_STD_HEAD = ("국가", "업체명", "연락처", "이메일")
_JUNK = re.compile(r"^[0-9?]+$")

# 파일명 토큰 → ISO2 (레지스트리 미반영 라이브 코드 대비 보강 — 표준 사실 매핑).
_NAME_ISO = {
    "대한민국": "KR", "한국": "KR", "멕시코": "MX", "싱가포르": "SG", "덴마크": "DK",
    "일본": "JP", "미국": "US", "독일": "DE", "영국": "GB", "호주": "AU",
}


def _country_from(text_: str) -> str | None:
    for tok, iso in _NAME_ISO.items():
        if tok in text_:
            return iso
    return None


def _norm(v) -> str:  # noqa: ANN001
    return str(v).strip() if v is not None else ""


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=r"C:\Users\WSCOPY\Desktop\DB")
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    sm = get_sessionmaker(Settings())
    st = {"filled_d": 0, "filled_c": 0, "no_signal": 0, "no_row": 0}
    with sm() as s:
        for path in sorted(glob.glob(os.path.join(args.dir, "*.xls[xm]"))):
            fname_iso = _country_from(os.path.basename(path))
            wb = load_workbook(path, read_only=True)
            try:
                for ws in wb.worksheets:
                    iso = fname_iso or _country_from(ws.title)
                    if iso is None:
                        continue
                    rows = ws.iter_rows(values_only=True)
                    try:
                        head = [_norm(c) for c in next(rows)]
                    except StopIteration:
                        continue
                    # 헤더 탐지형 — 표준 12컬럼뿐 아니라 청약공모('회사명'만)·코넥스
                    # ('회사명'+'홈페이지') 같은 비표준 서식도 처리(빈 국가 유입의 주범).
                    def col(*names):  # noqa: ANN202
                        for i, h in enumerate(head):
                            if h in names:
                                return i
                        return None
                    i_name = col("업체명", "회사명")
                    if i_name is None:
                        continue
                    i_site = col("사이트", "홈페이지", "site", "도메인")
                    i_country = col("국가")
                    blanks = 0
                    for r in rows:
                        vals = [_norm(c) for c in r]
                        if not any(vals):
                            blanks += 1
                            if blanks > 200:
                                break
                            continue
                        blanks = 0
                        get = lambda i: vals[i] if i is not None and i < len(vals) else ""  # noqa: E731
                        country_cell, name, site = get(i_country), get(i_name), get(i_site)
                        if country_cell and not _JUNK.match(country_cell):
                            continue  # 셀에 국가 있음 — 이미 정상 유입.
                        dom = normalize_domain(site)
                        if not dom and not name:
                            continue
                        keys = []
                        # 원 시드(import-existing)는 국가 '셀' 그대로 키를 만들어 빈 국가면
                        # 'name::…' 키다. fastpath 는 'KR' 폴백 키 — 둘 다 시도한다.
                        for c in ("", "KR"):
                            try:
                                k = canonical_key(domain=dom, name=name, country=c)
                            except ValueError:
                                continue
                            if k not in keys:
                                keys.append(k)
                        n1 = n2 = 0
                        for k in keys:
                            n1 += s.execute(text(
                                "update discovered_company set country=:c"
                                " where canonical_key=:k and (coalesce(country,'')=''"
                                " or country ~ '^[0-9?]+$')"
                            ), {"c": iso, "k": k}).rowcount
                            n2 += s.execute(text(
                                "update company set country=:c where canonical_key=:k"
                                " and (coalesce(country,'')='' or country ~ '^[0-9?]+$')"
                            ), {"c": iso, "k": k}).rowcount
                        st["filled_d"] += n1
                        st["filled_c"] += n2
                        if n1 == 0 and n2 == 0:
                            st["no_row"] += 1
            finally:
                wb.close()
        if args.apply:
            s.commit()
            print("[country-fname] COMMIT")
        else:
            s.rollback()
            print("[country-fname] DRY-RUN(롤백)")
    print(f"[country-fname] {st}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
