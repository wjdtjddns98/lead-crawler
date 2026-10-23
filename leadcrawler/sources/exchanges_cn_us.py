"""거래소 상장목록 발견 소스(Tier B) — CN(상하이 SSE·선전 SZSE) / US(Nasdaq 스크리너: NASDAQ·NYSE·AMEX).

2026-10-23 신설(exchanges_major 후속). 원장 실측 CN 상장 1,165(전수 ≈5,200)·US 2,727(EDGAR,
전수 ≈5,300) 로 큰 사각이 남아 있었다. 전부 공개 정적 HTTP.

라이브 엔드포인트(2026-10-23 이 PC 실측, 전부 200):
- SSE: ``query.sse.com.cn/sseQuery/commonQuery.do?sqlId=COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L
  &STOCK_TYPE=…`` (Referer 필수) — STOCK_TYPE 1(메인보드 1,700)·8(커촹반 618).
  ``pageHelp.pageSize=5000`` 1페이지.
  필드: A_STOCK_CODE·FULL_NAME·FULL_NAME_IN_ENGLISH·DELIST_DATE('-'=상장중). 웹사이트 없음 → Yahoo
  ``.SS``(yahoo_symbol 이 6/9 시작 코드로 판정).
- SZSE: ``szse.cn/api/report/ShowReport?SHOWTYPE=xlsx&CATALOGID=1110&TABKEY=tab1`` xlsx(Referer 필수)
  2,906행 — 公司全称·英文名称·A股代码·B股代码·所属行业·**公司网址 동봉**(도메인 해석 불필요).
- Nasdaq 스크리너: ``api.nasdaq.com/api/screener/stocks?tableonly=true&limit=10000&download=true
  &exchange={nasdaq|nyse|amex}`` (브라우저 UA 필수) — 전 미국 상장 6,999행(country 포함). ``country ==
  "United States"`` 5,271 중 워런트·유닛·우선주·권리·채권·예탁증서·ETF/펀드 등 비(非)보통주 종목을 이름
  패턴으로 제외하고, 종목명 꼬리("Common Stock"·"Class A Ordinary Shares"…)를 떼 법인명을 만든다.
  같은 법인의 복수 클래스(GOOG/GOOGL)는 법인명 기준 첫 심볼만. EDGAR(reg:sec) 와 키가 달라 원장에 둘 다
  남을 수 있지만, Yahoo 해석 후 같은 도메인이면 파이프라인 도메인 동치 dedup(``dedup.skip.resolved``)이
  건너뛴다 — 도메인 미해석 EDGAR 잔여행과의 이중은 수용(승격 시 company PK 가 canonical_key 결정적).

dry_run: 네트워크 없는 결정적 더미(베이스 ``_dry`` 상속).
"""

from __future__ import annotations

import io
import re
from typing import Any

import openpyxl

from ..dedup import normalize_domain
from ..logging import get_logger
from .base import DiscoveredCompany, Segment, build_company, english_display, opt_str
from .exchanges import ExchangeSource
from .exchanges_major import _aliases

log = get_logger("sources.exchanges_cn_us")

_BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/130.0 Safari/537.36"
)


# --- SSE(상하이) --------------------------------------------------------------

_SSE_URL = (
    "http://query.sse.com.cn/sseQuery/commonQuery.do?sqlId=COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L"
    "&STOCK_TYPE={stock_type}&REG_PROVINCE=&CSRC_CODE=&STOCK_CODE=&COMPANY_STATUS=2,4,5,7,8"
    "&type=inParams&isPagination=true&pageHelp.cacheSize=1&pageHelp.beginPage=1"
    "&pageHelp.pageSize=5000&pageHelp.pageNo=1"
)
_SSE_HEADERS = {"User-Agent": _BROWSER_UA, "Referer": "http://www.sse.com.cn/", "Accept": "*/*"}
_SSE_BOARDS = (("1", "SSE Main Board"), ("8", "SSE STAR Market"))


def _sse_rows(payload: Any) -> list[dict[str, Any]]:
    """SSE 응답 ``pageHelp.data`` → 상장중(DELIST_DATE '-'/빈값) 행만 공통 dict 로."""
    data = (payload.get("pageHelp") or {}).get("data") if isinstance(payload, dict) else None
    out = []
    for row in data if isinstance(data, list) else []:
        if not isinstance(row, dict):
            continue
        code, name = opt_str(row.get("A_STOCK_CODE")), opt_str(row.get("FULL_NAME"))
        delist = opt_str(row.get("DELIST_DATE"))
        if code and name and delist in (None, "-"):
            out.append({
                "code": code, "name": name, "eng": opt_str(row.get("FULL_NAME_IN_ENGLISH")),
            })
    return out


class SseSource(ExchangeSource):
    """상하이 증권거래소(SSE) 상장목록 JSON — 메인보드 + 커촹반(STAR)."""

    name = "sse"
    registry = "sse"
    countries = _aliases("CN")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        out: list[DiscoveredCompany] = []
        seen: set[str] = set()
        for stock_type, market in _SSE_BOARDS:
            try:
                rows = _sse_rows(fetcher.get_json(
                    _SSE_URL.format(stock_type=stock_type), headers=_SSE_HEADERS
                ))
            except Exception as exc:  # noqa: BLE001 — 한 보드 실패는 부분 결과(재수집 기록 생략).
                log.info("sse.error", board=market, err_type=type(exc).__name__, err=str(exc))
                self._partial = True
                continue
            if not rows:  # 200 이지만 스키마 이상/0행 — 완료 기록 시 7일 누락 → 부분 결과.
                log.warning("sse.parse.empty", board=market)
                self._partial = True
                continue
            for row in rows:
                if row["code"] in seen:
                    continue
                seen.add(row["code"])
                name, name_eng = english_display(row["name"], row["eng"], "CN")
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=name, name_eng=name_eng, domain=None,
                    registry=self.registry, registry_id=row["code"], ticker=row["code"],
                    market=market, listed_verified=True,
                ))
                if len(out) >= cap:
                    log.info("sse.live", segment=segment.label, n=len(out))
                    return out
        log.info("sse.live", segment=segment.label, n=len(out))
        return out


# --- SZSE(선전) ---------------------------------------------------------------

_SZSE_URL = (
    "http://www.szse.cn/api/report/ShowReport?SHOWTYPE=xlsx&CATALOGID=1110&TABKEY=tab1"
)
_SZSE_HEADERS = {"User-Agent": _BROWSER_UA, "Referer": "http://www.szse.cn/", "Accept": "*/*"}


def _parse_szse_workbook(wb: Any) -> list[dict[str, Any]]:
    """SZSE 상장회사 목록 xlsx — 1행 헤더(板块/公司全称/英文名称/A股代码/B股代码/公司网址…)."""
    ws = wb[wb.sheetnames[0]]
    col: dict[str, int] | None = None
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(values_only=True):
        if col is None:
            if any(isinstance(c, str) and c.strip() == "公司全称" for c in row):
                # 헤더 셀 안의 공백("B股 简 称"·"地      区")은 제거해 키를 안정화.
                col = {re.sub(r"\s+", "", c): i for i, c in enumerate(row) if isinstance(c, str)}
            continue

        def _cell(key: str) -> str | None:
            i = col.get(key)  # noqa: B023 — col 은 헤더 발견 후 고정.
            if i is None or i >= len(row) or row[i] is None:  # noqa: B023
                return None
            return opt_str(str(row[i]))

        name = _cell("公司全称")
        code = _cell("A股代码") or _cell("B股代码")
        if not name or not code:
            continue
        out.append({
            "code": code, "name": name, "eng": _cell("英文名称"), "website": _cell("公司网址"),
            "board": _cell("板块"),
        })
    return out


class SzseSource(ExchangeSource):
    """선전 증권거래소(SZSE) 상장회사 목록 xlsx — 웹사이트 동봉."""

    name = "szse"
    registry = "szse"
    countries = _aliases("CN")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        try:
            data = fetcher.get_bytes(_SZSE_URL, headers=_SZSE_HEADERS)
            rows = _parse_szse_workbook(openpyxl.load_workbook(io.BytesIO(data), data_only=True))
        except Exception as exc:  # noqa: BLE001
            log.info("szse.error", err_type=type(exc).__name__, err=str(exc))
            return []
        if not rows:
            log.warning("szse.parse.empty")
        out = []
        seen: set[str] = set()
        for row in rows:
            if row["code"] in seen:
                continue
            seen.add(row["code"])
            name, name_eng = english_display(row["name"], row["eng"], "CN")
            out.append(build_company(
                source=self.name, segment=listed_seg, name=name, name_eng=name_eng,
                domain=normalize_domain(row["website"]),
                registry=self.registry, registry_id=row["code"], ticker=row["code"],
                market=f"SZSE {row['board']}" if row["board"] else "SZSE", listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("szse.live", segment=segment.label, n=len(out))
        return out


# --- Nasdaq 스크리너(미국 전 거래소) -----------------------------------------------

_NASDAQ_URL = (
    "https://api.nasdaq.com/api/screener/stocks?tableonly=true&limit=10000&offset=0&download=true"
    "&exchange={exchange}"
)
_NASDAQ_HEADERS = {"User-Agent": _BROWSER_UA, "Accept": "application/json, text/plain, */*",
                   "Accept-Language": "en-US,en;q=0.9"}
_NASDAQ_EXCHANGES = (("nasdaq", "NASDAQ"), ("nyse", "NYSE"), ("amex", "NYSE American"))
_NASDAQ_US = "United States"
# 보통주가 아닌 종목(워런트·유닛·우선주·권리·채권/채무증권·예탁증서·ETF/펀드/신탁상품) — 2026-10-23 US
# 5,271행 실측 접미 분포 기준. "… Trust Common Stock"(REIT 법인)은 보통주라 유지한다.
_NASDAQ_NON_COMMON = re.compile(
    r"\b(Warrants?|Units?|Preferred|Preference|Rights?|Notes?|Debentures?|Bonds?|Depositary"
    r"|ETF|ETN|Fund|Trust Preferred|Subordinated|Senior|Fixed[- ]to[- ]Floating|Perpetual"
    r"|Index|Portfolio|Closed[- ]End)\b",
    re.I,
)
_NASDAQ_FUND_TAIL = re.compile(r"\b(Trust|Fund)\s*(Inc\.?|Ltd\.?)?\s*$", re.I)  # 꼬리 Trust = 상품.
# SPAC(블랭크체크) — 사업·IR 담당이 없는 껍데기 법인(2026-10-23 NASDAQ 표본 앞머리 다수).
_NASDAQ_SPAC = re.compile(r"\bAcquisition\s+(Corp|Co|Company|Inc|Ltd|Holdings|Group)\b", re.I)
# 종목명 꼬리(주식 종류) — 법인명만 남긴다.
_NASDAQ_NAME_TAIL = re.compile(
    r"\s*(?:\(The\)\s*)?(?:-\s*)?(?:(?:Class|Series)\s+[A-Z0-9]+\s+)?"
    r"(?:Common\s+(?:Stock|Shares?)|Ordinary\s+Shares?|Capital\s+Stock"
    r"|Common\s+Shares?\s+of\s+Beneficial\s+Interest|Shares?\s+of\s+Beneficial\s+Interest"
    r"|Common\s+Units?|Limited\s+Partnership\s+Units?)\s*$",
    re.I,
)
_NASDAQ_THE_RE = re.compile(r"\s*\(The\)\s*", re.I)


def _nasdaq_company_name(raw: str) -> str | None:
    """스크리너 종목명 → 법인명(보통주가 아니면 None)."""
    name = (raw or "").strip()
    if not name or _NASDAQ_NON_COMMON.search(name) or _NASDAQ_SPAC.search(name):
        return None
    base = _NASDAQ_NAME_TAIL.sub("", name).strip()
    if not base or _NASDAQ_FUND_TAIL.search(base) and "Common" not in name:
        return None
    base = _NASDAQ_THE_RE.sub(" ", base).strip(" -")
    return base or None


def _nasdaq_rows(payload: Any) -> list[dict[str, Any]]:
    data = payload.get("data") if isinstance(payload, dict) else None
    rows = data.get("rows") if isinstance(data, dict) else None
    out = []
    for row in rows if isinstance(rows, list) else []:
        if not isinstance(row, dict) or opt_str(row.get("country")) != _NASDAQ_US:
            continue
        symbol = opt_str(row.get("symbol"))
        name = _nasdaq_company_name(str(row.get("name") or ""))
        if symbol and name:
            out.append({"symbol": symbol, "name": name})
    return out


class NasdaqScreenerSource(ExchangeSource):
    """Nasdaq 스크리너로 미국 3대 거래소(NASDAQ·NYSE·NYSE American) 상장 보통주 발행사."""

    name = "nasdaq"
    registry = "nasdaq"
    countries = _aliases("US")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        out: list[DiscoveredCompany] = []
        seen_symbol: set[str] = set()
        seen_name: set[str] = set()
        for path, market in _NASDAQ_EXCHANGES:
            try:
                rows = _nasdaq_rows(fetcher.get_json(
                    _NASDAQ_URL.format(exchange=path), headers=_NASDAQ_HEADERS
                ))
            except Exception as exc:  # noqa: BLE001 — 한 거래소 실패는 부분 결과(재수집 기록 생략).
                log.info("nasdaq.error", exchange=market, err_type=type(exc).__name__, err=str(exc))
                self._partial = True
                continue
            if not rows:  # 200 이지만 스키마 이상/0행 — 부분 결과.
                log.warning("nasdaq.parse.empty", exchange=market)
                self._partial = True
                continue
            for row in rows:
                key = row["name"].lower()
                if row["symbol"] in seen_symbol or key in seen_name:
                    continue  # 복수 클래스(GOOG/GOOGL) 는 첫 심볼만.
                seen_symbol.add(row["symbol"])
                seen_name.add(key)
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=row["name"], domain=None,
                    registry=self.registry, registry_id=row["symbol"], ticker=row["symbol"],
                    market=market, listed_verified=True,
                ))
                if len(out) >= cap:
                    log.info("nasdaq.live", segment=segment.label, n=len(out))
                    return out
        log.info("nasdaq.live", segment=segment.label, n=len(out))
        return out
