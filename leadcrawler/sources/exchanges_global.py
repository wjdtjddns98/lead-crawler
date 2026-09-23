"""거래소 상장목록 발견 소스(Tier B, 유럽) — DE/Euronext(FR·IT·NL·BE·PT·IE·NO)/ES/CH.

``exchanges.py``(동남아 전용)와 분리한 신규 파일(2026-09-23, PR-b) — 국가별 상장기업
전수를 공개 정적 HTTP(+일부 curl_cffi)로 받는다. 베이스(:class:`ExchangeSource`,
``exchanges.py``)의 상장 세그먼트 게이트·``_seg``(업종 무도장)·``_memo``(국가별 1회)·
``cursor_store`` 재수집 주기를 그대로 물려받는다.

라이브 엔드포인트 상태(2026-09-23 이 PC 실측, 요청 거래소당 ≤5회):
- Xetra(독일): 랜딩 페이지에서 ``Listed-companies.xlsx`` blob URL 을 정규식으로 찾아
  GET(200, 71KB) → openpyxl 로 Prime/General/Scale 3시트 파싱(Basic Board 제외, PO 기본값).
  blob 해시는 재발행 시 바뀔 수 있어 랜딩 파싱이 정식, 실패 시 모듈 상수 URL 로 폴백.
- Euronext(FR/IT/NL/BE/PT/IE/NO): ``live.euronext.com/en/pd_es/data/stocks`` GET JSON
  (200, CloudFront, 챌린지 없음). 국가→MIC 매핑으로 1콜에 해당국 전 시장을 받는다.
  ``iDisplayLength`` 상한·IT ``EXGM`` 허용 여부는 미확인(스모크로 확인).
- BME(스페인): ``apiweb.bolsasymercados.es/Market/v1/EQ/ListedCompanies`` GET JSON(200).
  SIBE + BMEGrowth 두 패스. SICAV/HedgeFunds/ETF 세그먼트 행 제외.
  ``shareName`` 은 약칭이라 ticker 로 쓰지 않는다(ticker=None, Yahoo 는 companyKey 폴백
  이 무의미하면 미스 — 웹사이트 필드가 있는 행만 직접 도메인).
- SIX(스위스): ``six-group.com/sheldon/equity_issuers/v1/equity_issuers.json`` GET
  (200, content-type text/plain 이지만 본문은 JSON → get_text + json.loads).
  ``country``/``lastListingDate`` 값 형식은 미확인 — 관대한 필터로 시작(스모크로 조정).

dry_run: 네트워크 없는 결정적 더미(베이스 ``_dry`` 상속).
"""

from __future__ import annotations

import html
import io
import json
import re
from datetime import datetime, timezone
from typing import Any

import openpyxl

from ..dedup import normalize_domain
from ..logging import get_logger
from .base import Segment, build_company
from .countries import resolve_country, supported_countries
from .exchanges import ExchangeSource

log = get_logger("sources.exchanges_global")


# --- Deutsche Börse(Xetra) ------------------------------------------------

_DE_BASE = "https://www.cashmarket.deutsche-boerse.com"
_DE_LANDING_URL = f"{_DE_BASE}/cash-en/Data-Tech/statistics/listed-companies"
# 실측 blob(2026-09-23) — 랜딩 파싱 실패 시 폴백. 해시는 재발행 시 바뀔 수 있다.
_DE_FALLBACK_XLSX = (
    f"{_DE_BASE}/resource/blob/67858/3db127ac20c83f2955d06a43eddafede/"
    "data/Listed-companies.xlsx"
)
_DE_BLOB_RE = re.compile(r'href="([^"]*Listed-companies\.xlsx)"', re.I)
# Basic Board(비규제 시장)는 PO 기본값으로 제외.
_DE_SHEETS = ("Prime Standard", "General Standard", "Scale")


def _de_blob_url(landing_html: str) -> str | None:
    """랜딩 페이지 HTML 에서 xlsx blob 절대 URL 을 찾는다(못 찾으면 None → 폴백 URL)."""
    m = _DE_BLOB_RE.search(landing_html or "")
    if not m:
        return None
    href = html.unescape(m.group(1))
    return href if href.startswith("http") else _DE_BASE + href


def _parse_de_workbook(wb: Any) -> list[dict[str, Any]]:
    """DE 상장목록 워크북에서 (ISIN, 심볼, 회사명, 시트명) 행을 뽑는다(Country=Germany 만).

    헤더 행은 "ISIN" 셀 위치로 탐색한다(Scale 시트는 안내 행이 앞에 있어 행 번호 고정 불가).
    """
    out: list[dict[str, Any]] = []
    for sheet_name in _DE_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        col_map: dict[str, int] | None = None
        for row in wb[sheet_name].iter_rows(values_only=True):
            if col_map is None:
                if any(isinstance(c, str) and c.strip() == "ISIN" for c in row):
                    col_map = {c.strip(): i for i, c in enumerate(row) if isinstance(c, str)}
                continue
            if not row or all(c is None for c in row):
                continue
            isin_i = col_map.get("ISIN")
            if isin_i is None or isin_i >= len(row) or not row[isin_i]:
                continue
            country_i = col_map.get("Country")
            country = row[country_i] if country_i is not None and country_i < len(row) else None
            if country != "Germany":
                continue

            def _cell(key: str) -> Any:
                i = col_map.get(key)  # noqa: B023 — col_map 은 시트 루프마다 재바인딩.
                return row[i] if i is not None and i < len(row) else None  # noqa: B023

            out.append({
                "isin": row[isin_i],
                "ticker": _cell("Trading Symbol"),
                "name": _cell("Company"),
                "sheet": sheet_name,
            })
    return out


class DeutscheBoerseSource(ExchangeSource):
    """독일 증권거래소(Xetra) 상장목록 — 랜딩 blob 파싱 + openpyxl(모듈 docstring 참조)."""

    name = "xetra"
    registry = "xetra"
    countries = frozenset({
        "de", "deu", "germany", "deutschland", "federal republic of germany", "독일",
    })

    def _live(self, segment: Segment) -> list[Any]:
        fetcher = self._client()
        cap = self._settings.discovery_max_per_source
        listed_seg = self._seg(segment)
        try:
            landing = fetcher.get_text(_DE_LANDING_URL)
        except Exception as exc:  # noqa: BLE001 — 랜딩 실패해도 폴백 URL 로 계속.
            log.info("xetra.landing.error", err_type=type(exc).__name__, err=str(exc))
            landing = ""
        blob_url = _de_blob_url(landing)
        if blob_url is None:
            log.info("xetra.blob.fallback", url=_DE_FALLBACK_XLSX)
            blob_url = _DE_FALLBACK_XLSX
        try:
            data = fetcher.get_bytes(blob_url)
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — 네트워크/형식 이상 → graceful 빈 결과.
            log.info("xetra.error", url=blob_url, err_type=type(exc).__name__, err=str(exc))
            return []

        out = []
        seen: set[str] = set()
        for row in _parse_de_workbook(wb):
            isin = str(row["isin"]).strip()
            name = str(row["name"] or "").strip()
            if not isin or not name or isin in seen:
                continue
            seen.add(isin)
            ticker = str(row["ticker"]).strip() if row["ticker"] else None
            out.append(build_company(
                source=self.name, segment=listed_seg, name=name, domain=None,
                registry=self.registry, registry_id=isin, ticker=ticker,
                market=row["sheet"], listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("xetra.live", segment=segment.label, n=len(out))
        return out


# --- Euronext(FR/IT/NL/BE/PT/IE/NO) ---------------------------------------

# 국가→MIC(Market Identifier Code) 매핑(2026-09-23 설계). IT EXGM 허용 여부·MIC 목록의
# 실제 유효성은 라이브 스모크로 확인(0행이면 로그만 남기고 계속 — graceful).
_EURONEXT_MIC_BY_ISO2: dict[str, tuple[str, ...]] = {
    "FR": ("XPAR", "ALXP"),
    "IT": ("MTAA",),
    "NL": ("XAMS",),
    "BE": ("XBRU",),
    "PT": ("XLIS",),
    "IE": ("XMSM", "XESM"),
    "NO": ("XOSL", "MERK"),
}
_EURONEXT_MAX_PAGES = 30
_EURONEXT_PAGE_SIZE = 100
_EURONEXT_NAME_RE = re.compile(r">([^<]+)</a>")
_EURONEXT_TAG_RE = re.compile(r"<[^>]+>")
_EURONEXT_TITLE_RE = re.compile(r'title=[\'"]([^\'"]+)[\'"]')


def _euronext_countries() -> frozenset[str]:
    """``countries.py`` 에 이미 등록된 별칭을 Euronext 7개국 ISO2 로 필터해 재사용한다."""
    aliases: set[str] = set()
    for c in supported_countries():
        if c.iso2 in _EURONEXT_MIC_BY_ISO2:
            aliases.update(c.aliases)
    return frozenset(aliases)


def _parse_euronext_row(row: Any) -> dict[str, Any] | None:
    """``aaData`` 한 행(list, 셀 일부는 HTML 문자열)에서 회사 식별 정보를 뽑는다."""
    if not isinstance(row, list) or len(row) < 4:
        return None
    name_cell = str(row[1] or "")
    isin = str(row[2] or "").strip()
    symbol = str(row[3] or "").strip()
    if not isin:
        return None
    m = _EURONEXT_NAME_RE.search(name_cell)
    raw_name = m.group(1) if m else _EURONEXT_TAG_RE.sub("", name_cell)
    name = html.unescape(raw_name).strip()
    if not name:
        return None
    market = None
    if len(row) > 4:
        mt = _EURONEXT_TITLE_RE.search(str(row[4] or ""))
        market = html.unescape(mt.group(1)).strip() if mt else None
    return {"isin": isin, "name": name, "symbol": symbol or None, "market": market}


class EuronextSource(ExchangeSource):
    """Euronext(파리·밀라노·암스테르담·브뤼셀·리스본·더블린·오슬로) 상장목록.

    registry 는 단일 ``euronext``(ISIN 이 전역 유일이라 시장별 registry 분리 불필요,
    시장 구분은 ``ticker``/``market`` 필드로). 국가별 MIC 매핑 1콜로 해당국 전 시장 수집.
    """

    name = "euronext"
    registry = "euronext"
    countries = _euronext_countries()
    list_url = "https://live.euronext.com/en/pd_es/data/stocks"

    def _live(self, segment: Segment) -> list[Any]:
        fetcher = self._client()
        cap = self._settings.discovery_max_per_source
        listed_seg = self._seg(segment)
        country = resolve_country(segment.country)
        mics = _EURONEXT_MIC_BY_ISO2.get(country.iso2) if country else None
        if not mics:
            return []
        mics_param = ",".join(mics)

        out = []
        seen: set[str] = set()
        start = 0
        page = 0
        total: int | None = None
        while len(out) < cap and page < _EURONEXT_MAX_PAGES and (
            total is None or start < total
        ):
            try:
                payload = fetcher.get_json(self.list_url, params={
                    "mics": mics_param,
                    "display_datapoints": "dp_stocks",
                    "display_filters": "df_stocks2",
                    "iDisplayStart": start,
                    "iDisplayLength": _EURONEXT_PAGE_SIZE,
                })
            except Exception as exc:  # noqa: BLE001 — 네트워크/형식 이상 → 부분 결과 보존.
                log.info(
                    "euronext.error", mics=mics_param,
                    err_type=type(exc).__name__, err=str(exc),
                )
                break
            if not isinstance(payload, dict):
                break
            total = payload.get("iTotalRecords") or 0
            rows = payload.get("aaData")
            if not isinstance(rows, list) or not rows:
                break
            before = len(seen)
            for raw in rows:
                parsed = _parse_euronext_row(raw)
                if not parsed or parsed["isin"] in seen:
                    continue
                seen.add(parsed["isin"])
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=parsed["name"], domain=None,
                    registry=self.registry, registry_id=parsed["isin"],
                    ticker=parsed["symbol"], market=parsed["market"] or mics_param,
                    listed_verified=True,
                ))
                if len(out) >= cap:
                    break
            if len(seen) == before:  # 이번 페이지에서 신규 ISIN 0건 → 동일 페이지 반복(무한루프 방지).
                break
            start += len(rows)
            page += 1
        log.info("euronext.live", segment=segment.label, n=len(out))
        return out


# --- BME(스페인) -----------------------------------------------------------

_BME_MAX_PAGES = 50
_BME_PAGE_SIZE = 100
# 비영업 상품 세그먼트/시스템 — 실존 기업이 아닌 투자기구·ETF 배제(제약②).
_BME_EXCLUDE_SEGMENTS = frozenset({"SICAV", "HedgeFunds"})


class BmeSource(ExchangeSource):
    """스페인 증권거래소(BME) 상장목록 — SIBE + BME Growth 두 패스(PO 기본값=중소 병행시장 포함)."""

    name = "bme"
    registry = "bme"
    countries = frozenset({"es", "esp", "spain", "españa", "스페인"})
    list_url = "https://apiweb.bolsasymercados.es/Market/v1/EQ/ListedCompanies"

    def _fetch_pass(
        self, fetcher: Any, extra_params: dict[str, str], market_label: str,
        cap: int, listed_seg: Segment, out: list[Any], seen: set[str],
    ) -> None:
        page = 0
        while len(out) < cap and page < _BME_MAX_PAGES:
            params: dict[str, Any] = {"page": page, "pageSize": _BME_PAGE_SIZE, **extra_params}
            try:
                payload = fetcher.get_json(self.list_url, params=params)
            except Exception as exc:  # noqa: BLE001 — 네트워크/형식 이상 → 부분 결과 보존.
                log.info(
                    "bme.error", pass_=market_label,
                    err_type=type(exc).__name__, err=str(exc),
                )
                return
            if not isinstance(payload, dict):
                return
            rows = payload.get("data")
            if not isinstance(rows, list) or not rows:
                return
            before = len(seen)
            for row in rows:
                if not isinstance(row, dict):
                    continue
                if (
                    row.get("mtfSegment") in _BME_EXCLUDE_SEGMENTS
                    or row.get("tradingSystem") == "ETF"
                ):
                    continue  # SICAV/HedgeFunds/ETF — 실존 영업기업 아님(제약②).
                key = str(row.get("companyKey") or "").strip()
                name = str(row.get("name") or "").strip()
                if not key or not name or key in seen:
                    continue
                seen.add(key)
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=name,
                    domain=normalize_domain(row.get("website")),
                    registry=self.registry, registry_id=key, ticker=None,
                    market=market_label, listed_verified=True,
                ))
                if len(out) >= cap:
                    return
            if len(seen) == before:  # 이번 페이지에서 신규 companyKey 0건 → 동일 페이지 반복.
                return
            if not payload.get("hasMoreResults"):
                return
            page += 1

    def _live(self, segment: Segment) -> list[Any]:
        fetcher = self._client()
        cap = self._settings.discovery_max_per_source
        listed_seg = self._seg(segment)
        out: list[Any] = []
        seen: set[str] = set()
        self._fetch_pass(fetcher, {"tradingSystem": "SIBE"}, "SIBE", cap, listed_seg, out, seen)
        self._fetch_pass(
            fetcher, {"mtfSegment": "BMEGrowth"}, "BME Growth", cap, listed_seg, out, seen,
        )
        log.info("bme.live", segment=segment.label, n=len(out))
        return out


# --- SIX(스위스) -----------------------------------------------------------

_SIX_LIST_URL = "https://www.six-group.com/sheldon/equity_issuers/v1/equity_issuers.json"
# 값 형식(country="CH" 대 "Switzerland")이 미확인이라 둘 다 관대히 인정(스모크로 확정).
_SIX_COUNTRY_OK = frozenset({"CH", "SWITZERLAND"})


def _six_today() -> int:
    return int(datetime.now(timezone.utc).strftime("%Y%m%d"))


def _six_row_active(row: Any, today: int) -> bool:
    """CH 1차 상장이고 상장폐지일이 아직(또는 영구, 99991231) 도래하지 않은 행인지."""
    if not isinstance(row, dict):
        return False
    if str(row.get("country") or "").strip().upper() not in _SIX_COUNTRY_OK:
        return False
    if row.get("primaryListing") is not True:
        return False
    try:
        last = int(row.get("lastListingDate"))
    except (TypeError, ValueError):
        return False
    return last >= today


class SixSource(ExchangeSource):
    """스위스 증권거래소(SIX) 상장목록 — JSON(content-type text/plain, get_text+json.loads)."""

    name = "six"
    registry = "six"
    countries = frozenset({"ch", "che", "switzerland", "schweiz", "suisse", "스위스"})
    list_url = _SIX_LIST_URL

    def _live(self, segment: Segment) -> list[Any]:
        fetcher = self._client()
        cap = self._settings.discovery_max_per_source
        listed_seg = self._seg(segment)
        try:
            payload = json.loads(fetcher.get_text(self.list_url))
        except Exception as exc:  # noqa: BLE001 — 네트워크/형식 이상 → graceful 빈 결과.
            log.info("six.error", err_type=type(exc).__name__, err=str(exc))
            return []
        items = payload.get("itemList") if isinstance(payload, dict) else None
        if not isinstance(items, list):
            return []

        out = []
        seen: set[str] = set()
        today = _six_today()
        for row in items:
            if not _six_row_active(row, today):
                continue
            valor = str(row.get("valorNumber") or "").strip()
            name = str(row.get("company") or "").strip()
            if not valor or not name or valor in seen:
                continue
            seen.add(valor)
            symbol = str(row.get("valorSymbol") or "").strip() or None
            out.append(build_company(
                source=self.name, segment=listed_seg, name=name, domain=None,
                registry=self.registry, registry_id=valor, ticker=symbol,
                market="SIX Swiss Exchange", listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("six.live", segment=segment.label, n=len(out))
        return out
