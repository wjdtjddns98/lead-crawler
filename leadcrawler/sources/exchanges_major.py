"""거래소 상장목록 발견 소스(Tier B) — AU(ASX)/IN(NSE)/CA(TSX·TSXV)/HK(HKEX)/TW(TWSE·TPEx).

2026-10-23 신설. 배경: 상장 S 잡이 72개국을 돌지만 이 5개국은 거래소 소스가 없어 발견이
TW·HK 0건, AU 360·IN 494·CA 181 에 그쳤다(원장 실측). 전부 공개 정적 HTTP(인증·WAF 없음)로
상장 전수를 받는다. 베이스(:class:`ExchangeSource`, ``exchanges.py``)의 상장 세그먼트 게이트·
``_seg``(업종 무도장)·``_memo``(국가별 1회)·``cursor_store`` 재수집 주기를 그대로 물려받는다.

라이브 엔드포인트(2026-10-23 이 PC 실측, 전부 200):
- ASX: ``asx.com.au/asx/research/ASXListedCompanies.csv`` — 1행 "as at" 문구·2행 공백·3행 헤더
  (Company name, ASX code, GICS industry group), 2,050행. 웹사이트 없음 → Yahoo ``.AX``.
- NSE: ``nsearchives.nseindia.com/content/equities/EQUITY_L.csv`` — SYMBOL/NAME OF COMPANY/
  SERIES/ISIN 등(헤더 일부에 선행 공백). EQ 2,359·BE 221·BZ 33. BZ(정지성 거래 구간)는 제외.
  웹사이트 없음 → Yahoo ``.NS``.
- TSX: ``tsx.com/json/company-directory/search/{tsx|tsxv}/^*`` — ``results[]`` 2,334(TSX)·1,415
  (TSXV). TSX 목록은 ETF·CDR·채권펀드·우선주가 절반을 넘어 이름 패턴으로 걸러낸다(REIT 는 유지).
  웹사이트 없음 → Yahoo ``.TO``(클래스/유닛 구분자 ``.`` 는 Yahoo 표기 ``-`` 로).
- HKEX: ``hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx`` —
  3행 헤더(Stock Code/Name of Securities/Category/Sub-Category/…/ISIN), Category ``Equity`` +
  Sub-Category ``Equity Securities (Main Board|GEM)`` 만. 웹사이트 없음 → Yahoo ``.HK``
  (5자리 코드 선행 0 제거 후 4자리로).
- TWSE/TPEx: ``openapi.twse.com.tw/v1/opendata/t187ap03_L``(상장 1,097)·
  ``tpex.org.tw/openapi/v1/mopsfin_t187ap03_O``(상궤 894) JSON — **웹사이트 필드 동봉**(網址 /
  WebAddress) 이라 도메인 해석 없이 승격 가능. TWSE 는 英文簡稱 으로 영문 표시명(#412).

dry_run: 네트워크 없는 결정적 더미(베이스 ``_dry`` 상속).
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

import openpyxl

from ..dedup import normalize_domain
from ..logging import get_logger
from .base import DiscoveredCompany, Segment, build_company, english_display, opt_str
from .countries import supported_countries
from .exchanges import ExchangeSource

log = get_logger("sources.exchanges_major")


def _aliases(*iso2: str) -> frozenset[str]:
    """``countries.py`` 에 등록된 별칭을 해당 ISO2 로 필터해 재사용한다(exchanges_global 관례)."""
    want = set(iso2)
    out: set[str] = set()
    for c in supported_countries():
        if c.iso2 in want:
            out.update(c.aliases)
    return frozenset(out)


# --- ASX(호주) --------------------------------------------------------------

_ASX_CSV_URL = "https://www.asx.com.au/asx/research/ASXListedCompanies.csv"


def _parse_asx_csv(text: str) -> list[dict[str, str]]:
    """ASX CSV — 헤더 행("Company name" 으로 시작)을 찾아 그 아래를 DictReader 로 읽는다."""
    lines = (text or "").lstrip("﻿").splitlines()
    start = next(
        (i for i, ln in enumerate(lines) if ln.strip('"').startswith("Company name")), None
    )
    if start is None:
        return []
    out = []
    for row in csv.DictReader(io.StringIO("\n".join(lines[start:]))):
        code = opt_str(row.get("ASX code"))
        name = opt_str(row.get("Company name"))
        if code and name:
            out.append({"code": code, "name": name})
    if not out:
        log.warning("asx.parse.empty", lines=len(lines))
    return out


class AsxSource(ExchangeSource):
    """호주 증권거래소(ASX) 상장목록 CSV."""

    name = "asx"
    registry = "asx"
    countries = _aliases("AU")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        try:
            rows = _parse_asx_csv(fetcher.get_text(_ASX_CSV_URL))
        except Exception as exc:  # noqa: BLE001 — 네트워크/형식 이상 → graceful 빈 결과.
            log.info("asx.error", err_type=type(exc).__name__, err=str(exc))
            return []
        out = []
        seen: set[str] = set()
        for row in rows:
            if row["code"] in seen:
                continue
            seen.add(row["code"])
            out.append(build_company(
                source=self.name, segment=listed_seg, name=row["name"], domain=None,
                registry=self.registry, registry_id=row["code"], ticker=row["code"],
                market="ASX", listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("asx.live", segment=segment.label, n=len(out))
        return out


# --- NSE(인도) --------------------------------------------------------------

_NSE_CSV_URL = "https://nsearchives.nseindia.com/content/equities/EQUITY_L.csv"
_NSE_SERIES = frozenset({"EQ", "BE"})  # BZ(정지성 T+T 구간) 제외.
# nsearchives 는 비브라우저 UA(LeadCrawler/1.0) 요청을 응답 없이 끌어 타임아웃시킨다(2026-10-23 실측
# — Chrome UA 는 0.8s/200). 정적 CSV 1건이라 브라우저 UA 로 보낸다.
_NSE_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/130.0 Safari/537.36"
    ),
    "Accept": "text/csv,*/*",
}


def _parse_nse_csv(text: str) -> list[dict[str, str]]:
    """NSE EQUITY_L.csv — 헤더 선행 공백 정규화 후 EQ/BE 시리즈 행만."""
    reader = csv.reader(io.StringIO((text or "").lstrip("﻿")))
    header = next(reader, None)
    if not header:
        return []
    idx = {h.strip().upper(): i for i, h in enumerate(header)}
    need = ("SYMBOL", "NAME OF COMPANY", "SERIES", "ISIN NUMBER")
    if any(k not in idx for k in need):
        log.warning("nse.parse.no_header", header=header[:8])
        return []
    out = []
    for row in reader:
        if len(row) <= max(idx[k] for k in need):
            continue
        series = row[idx["SERIES"]].strip().upper()
        symbol = row[idx["SYMBOL"]].strip()
        name = row[idx["NAME OF COMPANY"]].strip()
        isin = row[idx["ISIN NUMBER"]].strip()
        if series in _NSE_SERIES and symbol and name and isin:
            out.append({"symbol": symbol, "name": name, "isin": isin})
    return out


def _nse_issuer_id(isin: str) -> str:
    """인도 ISIN(INE + 발행사 4자리 + 증권 2자리 + 일련 2자리 + 검증 1자리)에서 **발행사 7자리**만.

    심볼은 사명 변경 때, 전체 ISIN 은 액면분할·병합 때(일련 01→02) 바뀌어 같은 회사가 새
    canonical_key 로 재추출될 위험(제약①) — 발행사 부분은 유지된다(리뷰 MED, 2026-10-23).
    """
    return isin[:7].upper()


class NseSource(ExchangeSource):
    """인도 국립증권거래소(NSE) 상장목록 CSV. registry_id=ISIN 발행사 7자리, ticker=SYMBOL."""

    name = "nse"
    registry = "nse"
    countries = _aliases("IN")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        try:
            rows = _parse_nse_csv(fetcher.get_text(_NSE_CSV_URL, headers=_NSE_HEADERS))
        except Exception as exc:  # noqa: BLE001
            log.info("nse.error", err_type=type(exc).__name__, err=str(exc))
            return []
        out = []
        seen: set[str] = set()
        for row in rows:
            issuer = _nse_issuer_id(row["isin"])
            if issuer in seen:
                continue  # 같은 발행사의 복수 증권은 첫 행만.
            seen.add(issuer)
            out.append(build_company(
                source=self.name, segment=listed_seg, name=row["name"], domain=None,
                registry=self.registry, registry_id=issuer, ticker=row["symbol"],
                market="NSE", listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("nse.live", segment=segment.label, n=len(out))
        return out


# --- TSX·TSXV(캐나다) --------------------------------------------------------

_TSX_URL = "https://www.tsx.com/json/company-directory/search/{market}/^*"
_TSX_MARKETS = (("tsx", "TSX"), ("tsxv", "TSXV"))
# 운영 기업이 아닌 상품(ETF·CDR 예탁증서·채권/펀드·우선주·분할주·채무증권) — 2026-10-23 TSX 목록
# 2,334 중 1,623 이 해당. REIT/Trust 는 부동산 운영 법인이라 유지.
_TSX_NON_COMPANY = re.compile(
    r"\b(ETF|CDR|Index|Bond|Fund|Portfolio|Split|Preferred|Pref|Debenture|Note)s?\b", re.I
)


def _tsx_is_company(name: str) -> bool:
    return not _TSX_NON_COMPANY.search(name or "")


def _tsx_yahoo_ticker(symbol: str) -> str:
    """TSX 심볼의 클래스/유닛 구분자(``BBD.B``·``REI.UN``)는 Yahoo 표기 ``-`` 로."""
    return symbol.replace(".", "-")


class TsxSource(ExchangeSource):
    """토론토 증권거래소(TSX) + 벤처(TSXV) 회사 디렉터리 JSON."""

    name = "tsx"
    registry = "tsx"
    countries = _aliases("CA")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        out: list[DiscoveredCompany] = []
        seen: set[str] = set()
        for path, market in _TSX_MARKETS:
            try:
                payload = fetcher.get_json(_TSX_URL.format(market=path))
            except Exception as exc:  # noqa: BLE001 — 한 시장 실패는 부분 결과(재수집 기록 생략).
                log.info("tsx.error", market=market, err_type=type(exc).__name__, err=str(exc))
                self._partial = True
                continue
            results = payload.get("results") if isinstance(payload, dict) else None
            if not isinstance(results, list) or not results:
                # 200 이지만 스키마 이상/0행 — 완료로 기록하면 7일간 누락되므로 부분 결과 처리.
                log.warning("tsx.parse.empty", market=market)
                self._partial = True
                continue
            for row in results:
                if not isinstance(row, dict):
                    continue
                symbol = opt_str(row.get("symbol"))
                name = opt_str(row.get("name"))
                if not symbol or not name or symbol in seen or not _tsx_is_company(name):
                    continue
                seen.add(symbol)
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=name, domain=None,
                    registry=self.registry, registry_id=symbol, ticker=_tsx_yahoo_ticker(symbol),
                    market=market, listed_verified=True,
                ))
                if len(out) >= cap:
                    log.info("tsx.live", segment=segment.label, n=len(out))
                    return out
        log.info("tsx.live", segment=segment.label, n=len(out))
        return out


# --- HKEX(홍콩) --------------------------------------------------------------

_HKEX_XLSX_URL = (
    "https://www.hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx"
)
_HKEX_EQUITY_SUB = "Equity Securities"


def _parse_hkex_workbook(wb: Any) -> list[dict[str, str | None]]:
    """HKEX 증권 목록 워크북 — 헤더("Stock Code" 셀) 아래 Equity Securities 행만."""
    ws = wb[wb.sheetnames[0]]
    col: dict[str, int] | None = None
    out: list[dict[str, str | None]] = []
    for row in ws.iter_rows(values_only=True):
        if col is None:
            if any(isinstance(c, str) and c.strip() == "Stock Code" for c in row):
                col = {c.strip(): i for i, c in enumerate(row) if isinstance(c, str)}
            continue
        if not row:
            continue

        def _cell(key: str) -> str | None:
            i = col.get(key)  # noqa: B023 — col 은 헤더 발견 후 고정.
            if i is None or i >= len(row) or row[i] is None:
                return None
            return opt_str(str(row[i]))

        code, name = _cell("Stock Code"), _cell("Name of Securities")
        sub = _cell("Sub-Category") or ""
        if not code or not name or _cell("Category") != "Equity":
            continue
        if not sub.startswith(_HKEX_EQUITY_SUB):
            continue
        # HKD-RMB 듀얼카운터의 RMB 쪽(80700 TENCENT-R 등 24행, 2026-10-23 실측)은 같은 회사의
        # 중복 — Trading Currency 열이 있으면 HKD 만(열이 없는 구형식이면 통과).
        currency = _cell("Trading Currency")
        if currency is not None and currency.upper() != "HKD":
            continue
        code = code.zfill(5)  # 숫자 셀로 와도 5자리 문자열로 안정화.
        out.append({"code": code, "name": name, "sub": sub, "isin": _cell("ISIN")})
    return out


def _hkex_market(sub: str) -> str:
    m = re.search(r"\(([^)]+)\)", sub or "")
    return f"HKEX {m.group(1)}" if m else "HKEX"


class HkexSource(ExchangeSource):
    """홍콩 거래소(HKEX) 증권 목록 xlsx — 주식(Main Board·GEM)만."""

    name = "hkex"
    registry = "hkex"
    countries = _aliases("HK")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        try:
            data = fetcher.get_bytes(_HKEX_XLSX_URL)
            # read_only 모드는 이 파일의 치수 메타가 틀려 3~4행만 돌려준다(2026-10-23 실측) → 일반 로드.
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            rows = _parse_hkex_workbook(wb)
        except Exception as exc:  # noqa: BLE001
            log.info("hkex.error", err_type=type(exc).__name__, err=str(exc))
            return []
        out = []
        seen: set[str] = set()
        for row in rows:
            code = str(row["code"])
            if code in seen:
                continue
            seen.add(code)
            # Yahoo 는 4자리("0005.HK") — 5자리 코드의 선행 0 을 떼고 yahoo_symbol 이 zfill(4).
            ticker = code.lstrip("0") or "0"
            out.append(build_company(
                source=self.name, segment=listed_seg, name=str(row["name"]), domain=None,
                registry=self.registry, registry_id=code, ticker=ticker,
                market=_hkex_market(str(row["sub"])), listed_verified=True,
            ))
            if len(out) >= cap:
                break
        log.info("hkex.live", segment=segment.label, n=len(out))
        return out


# --- TWSE·TPEx(대만) ----------------------------------------------------------

_TWSE_URL = "https://openapi.twse.com.tw/v1/opendata/t187ap03_L"
_TPEX_URL = "https://www.tpex.org.tw/openapi/v1/mopsfin_t187ap03_O"


def _twse_rows(payload: Any) -> list[dict[str, Any]]:
    """TWSE(중문 키) 응답 → 공통 dict(code/name/eng/website/market)."""
    out = []
    for row in payload if isinstance(payload, list) else []:
        if not isinstance(row, dict):
            continue
        code, name = opt_str(row.get("公司代號")), opt_str(row.get("公司名稱"))
        if code and name:
            out.append({"code": code, "name": name, "eng": opt_str(row.get("英文簡稱")),
                        "website": opt_str(row.get("網址")), "market": "TWSE"})
    return out


def _tpex_rows(payload: Any) -> list[dict[str, Any]]:
    """TPEx(영문 키) 응답 → 공통 dict. 영문명 필드는 없다."""
    out = []
    for row in payload if isinstance(payload, list) else []:
        if not isinstance(row, dict):
            continue
        code, name = opt_str(row.get("SecuritiesCompanyCode")), opt_str(row.get("CompanyName"))
        if code and name:
            out.append({"code": code, "name": name, "eng": None,
                        "website": opt_str(row.get("WebAddress")), "market": "TPEx"})
    return out


class TwseSource(ExchangeSource):
    """대만 증권거래소(TWSE 상장) + 櫃買센터(TPEx 상궤) 공개 API — 웹사이트 동봉."""

    name = "twse"
    registry = "twse"
    countries = _aliases("TW")

    def _live(self, segment: Segment) -> list[DiscoveredCompany]:
        fetcher = self._client()
        cap = self._settings.exchange_max_per_source
        listed_seg = self._seg(segment)
        out: list[DiscoveredCompany] = []
        seen: set[str] = set()
        for url, parse in ((_TWSE_URL, _twse_rows), (_TPEX_URL, _tpex_rows)):
            try:
                rows = parse(fetcher.get_json(url))
            except Exception as exc:  # noqa: BLE001 — 한 시장 실패는 부분 결과(재수집 기록 생략).
                log.info("twse.error", url=url, err_type=type(exc).__name__, err=str(exc))
                self._partial = True
                continue
            if not rows:  # 200 이지만 스키마 이상/0행 — 완료 기록 시 7일 누락 → 부분 결과.
                log.warning("twse.parse.empty", url=url)
                self._partial = True
                continue
            for row in rows:
                if row["code"] in seen:
                    continue
                seen.add(row["code"])
                name, name_eng = english_display(row["name"], row["eng"], "TW")
                out.append(build_company(
                    source=self.name, segment=listed_seg, name=name, name_eng=name_eng,
                    domain=normalize_domain(row["website"]),
                    registry=self.registry, registry_id=row["code"], ticker=row["code"],
                    market=row["market"], listed_verified=True,
                ))
                if len(out) >= cap:
                    log.info("twse.live", segment=segment.label, n=len(out))
                    return out
        log.info("twse.live", segment=segment.label, n=len(out))
        return out
